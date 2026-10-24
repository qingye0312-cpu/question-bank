from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path(__file__).resolve().parents[1]
out=root/'templates';out.mkdir(exist_ok=True)
doc=Document();section=doc.sections[0];section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=section.bottom_margin=Cm(1.8);section.left_margin=section.right_margin=Cm(2)
for name in ('Normal','Title','Heading 1','Heading 2'):
 s=doc.styles[name];s.font.name='Heiti SC';s.font.color.rgb=RGBColor(0,0,0)
 fonts=s._element.get_or_add_rPr().rFonts
 for attr in list(fonts.attrib):
  if attr.endswith('Theme'):del fonts.attrib[attr]
 fonts.set(qn('w:eastAsia'),'Heiti SC')
for border in doc.styles.element.xpath('.//w:pBdr'):
 border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(15)
def table(rows):
 t=doc.add_table(rows=0,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;t.columns[0].width=Cm(3.1);t.columns[1].width=Cm(13.8)
 borders=OxmlElement('w:tblBorders')
 for tag in ['top','left','bottom','right','insideH','insideV']:
  edge=OxmlElement('w:'+tag);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
 t._tbl.tblPr.append(borders)
 for key,value in rows:
  cells=t.add_row().cells
  for j,text in enumerate([key,value]):
   c=cells[j];c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.width=Cm(3.1 if j==0 else 13.8);c.text=text
   for p in c.paragraphs:
    p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
    for r in p.runs:r.font.size=Pt(9.5);r.bold=j==0
   margins=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'65' if side in ('top','bottom') else '110');el.set(qn('w:type'),'dxa');margins.append(el)
   c._tc.get_or_add_tcPr().append(margins)
   if j==0:shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EDF1F3');c._tc.get_or_add_tcPr().append(shd)
  tr=t.rows[-1]._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
 return t
doc.add_paragraph('题库标准录入模板',style='Title')
doc.add_paragraph('模板版本 1.0')
doc.add_paragraph('使用本模板填写题目后，在题库程序中选择“导入”。每道题使用一个独立的两列表格，左侧字段名称保持不变，右侧填写内容。')
table([('模板版本','1.0'),('默认题库','模板示例')])
doc.add_paragraph('填写方法',style='Heading 1')
for text in ['后续三页分别为单选、多选、判断示例。开始录题时替换示例内容，或复制整张题目表继续填写；删除未使用的整张示例表，避免误导入。','选择题填写 2 至 8 个连续选项。单选答案填 B，多选填 A|C。判断题答案填“正确”或“错误”，选项全部留空。','新题的系统ID和基准修订号留空。修改系统导出的题目时，必须保留这两个字段，避免覆盖较新的内容。','正文、代码可保留换行。不要用颜色、加粗、勾号或自动编号表达答案。不能合并字段单元格，也不能插入嵌套表格。','本版支持文本与代码。图片、原生公式、修订和嵌入对象会在导入时报告，请保留原文件。','导入后先查看预览中的题目数量和内容，再确认写入。录题模板只规定格式，正确答案仍需核对。']:
 doc.add_paragraph(text)
examples=json.loads((root/'src-tauri/resources/examples.json').read_text())
headers=['外部编号','题型','题库','章节','题干',*['选项'+k for k in 'ABCDEFGH'],'答案','解析','标签','难度','适用版本','来源说明','系统ID','基准修订号']
for q in examples:
 doc.add_page_break();kind={'single':'单选','multiple':'多选','true_false':'判断','programming':'编程'}[q['type']]
 doc.add_paragraph(kind+'题示例',style='Heading 1');doc.add_paragraph('替换右列内容后使用，左列字段名保持不变。')
 vals=[q['external_id'],kind,'模板示例','数学基础',q['stem'],*[next((o['content'] for o in q['options'] if o['key']==k),'') for k in 'ABCDEFGH'],('正确' if q['answer'] else '错误') if isinstance(q['answer'],bool) else '|'.join(q['answer']),q['explanation'],'基础运算','简单','','示例题目','','']
 table(list(zip(headers,vals)))
doc.core_properties.title='题库标准录入模板';doc.core_properties.author='';doc.core_properties.subject='schema_version 1.0'
doc.save(out/'题库导入模板_v1.docx')
print(out/'题库导入模板_v1.docx')
